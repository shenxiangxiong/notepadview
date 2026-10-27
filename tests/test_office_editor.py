"""Office 编辑器的独立行为测试，不启动主窗口。"""
import base64
import io
import json
import zipfile
from datetime import datetime

import pytest
from PySide6.QtGui import QColor, QFont, QTextCharFormat, QTextCursor
from test_ui import app
from docx import Document
from docx.shared import Pt, RGBColor
from openpyxl import Workbook, load_workbook


def editor_class():
    from notepadview.office_editor import OfficeEditor
    return OfficeEditor


def test_blank_word_undo_snapshot_and_save(app, tmp_path):
    editor = editor_class().blank('docx')
    editor.text_edit.insertPlainText('正文')
    assert editor.undo_stack.canUndo()
    restored = editor_class().from_snapshot(json.loads(json.dumps(editor.snapshot())))
    assert restored.text_edit.toPlainText() == '正文'
    restored.undo_stack.undo()
    assert restored.text_edit.toPlainText() == ''
    restored.undo_stack.redo()
    target = tmp_path / 'new.docx'
    restored.save(target)
    assert Document(target).paragraphs[0].text == '正文'
    assert restored.undo_stack.isClean()


def test_word_format_table_and_image_roundtrip(app, tmp_path):
    source = tmp_path / 'source.docx'
    doc = Document()
    run = doc.add_paragraph().add_run('格式')
    run.bold, run.italic, run.underline = True, True, True
    run.font.name, run.font.size = 'Arial', Pt(16)
    run.font.color.rgb = RGBColor(20, 40, 60)
    doc.add_table(rows=2, cols=2).cell(0, 0).text = '表格'
    # 使用 Qt 生成有效 PNG，避免不同 libpng 的 CRC 容忍差异。
    from PySide6.QtCore import QBuffer, QIODevice
    from PySide6.QtGui import QImage
    image = QImage(4, 4, QImage.Format_RGB32)
    image.fill(QColor('red'))
    buffer = QBuffer()
    buffer.open(QIODevice.WriteOnly)
    image.save(buffer, 'PNG')
    doc.add_picture(io.BytesIO(bytes(buffer.data())))
    doc.save(source)
    editor = editor_class().from_path(source)
    assert '表格' in editor.text_edit.toPlainText()
    target = tmp_path / 'roundtrip.docx'
    editor.save(target)
    result = Document(target)
    run = next(r for p in result.paragraphs for r in p.runs if r.text == '格式')
    assert run.bold and run.italic and run.underline
    assert run.font.size.pt == 16
    assert run.font.color.rgb == RGBColor(20, 40, 60)
    assert result.tables[0].cell(0, 0).text == '表格'
    assert len(result.inline_shapes) == 1


def test_excel_types_formula_sheets_and_snapshot(app, tmp_path):
    source = tmp_path / 'source.xlsx'
    book = Workbook()
    book.active.append([12, True, datetime(2025, 1, 2), '=A1*2'])
    book.active['A1'].font = __import__('openpyxl').styles.Font(bold=True)
    book.save(source)
    editor = editor_class().from_path(source)
    editor.table.item(0, 0).setText('42')
    editor.add_sheet('第二页')
    editor.rename_sheet('改名')
    restored = editor_class().from_snapshot(json.loads(json.dumps(editor.snapshot())))
    restored.undo_stack.undo()
    assert restored.workbook.sheetnames[-1] == '第二页'
    restored.undo_stack.redo()
    target = tmp_path / 'out.xlsx'
    restored.save(target)
    result = load_workbook(target, data_only=False)
    assert result.active['A1'].value == 42
    assert result.active['A1'].font.bold
    assert result.active['B1'].value is True
    assert isinstance(result.active['C1'].value, datetime)
    assert result.active['D1'].value == '=A1*2'
    assert result.sheetnames[-1] == '改名'
    restored.delete_sheet()
    restored.delete_sheet()
    assert len(restored.workbook.sheetnames) == 1


@pytest.mark.parametrize('kind,part', [('docx', 'word/footnotes.xml'), ('xlsx', 'xl/vbaProject.bin'), ('xlsx', 'xl/externalLinks/externalLink1.xml')])
def test_lossy_packages_can_be_viewed_but_not_saved(app, tmp_path, kind, part):
    source = tmp_path / ('advanced.' + kind)
    editor_class().blank(kind).save(source)
    with zipfile.ZipFile(source, 'a') as archive:
        archive.writestr(part, '<advanced/>')
    editor = editor_class().from_path(source)
    with pytest.raises(ValueError):
        editor.save(tmp_path / ('unsafe.' + kind))


@pytest.mark.parametrize('kind', ['docx', 'xlsx'])
def test_atomic_failure_keeps_file_and_dirty_state(app, tmp_path, monkeypatch, kind):
    import notepadview.office_editor as module
    editor = editor_class().blank(kind)
    if kind == 'docx':
        editor.text_edit.insertPlainText('未保存')
    else:
        editor.table.item(0, 0).setText('23')
    target = tmp_path / ('existing.' + kind)
    target.write_bytes(b'original')
    monkeypatch.setattr(module.os, 'replace', lambda *args: (_ for _ in ()).throw(OSError('denied')))
    with pytest.raises(OSError):
        editor.save(target)
    assert target.read_bytes() == b'original'
    assert not editor.undo_stack.isClean()
    assert list(tmp_path.iterdir()) == [target]


def test_word_edits_preserve_format_and_clean_history(app, tmp_path):
    editor = editor_class().blank('docx')
    editor.text_edit.insertPlainText('before')
    editor.save(tmp_path / 'clean.docx')
    editor.text_edit.selectAll()
    editor._format('bold', True)
    restored = editor_class().from_snapshot(editor.snapshot())
    assert not restored.undo_stack.isClean()
    restored.undo_stack.undo()
    assert restored.undo_stack.isClean()
    restored.undo_stack.redo()
    restored.save(tmp_path / 'bold.docx')
    assert Document(tmp_path / 'bold.docx').paragraphs[0].runs[0].bold


def test_rejects_zip_bomb_and_xml_entities(app, tmp_path):
    path = tmp_path / 'bomb.docx'
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('word/document.xml', b'0' * 2000000)
    with pytest.raises(ValueError):
        editor_class().from_path(path)
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('word/document.xml', '<!DOCTYPE x [<!ENTITY y "z">]><x/>')
    with pytest.raises(ValueError):
        editor_class().from_path(path)


def test_excel_merged_cells_and_formula_text(app, tmp_path):
    path = tmp_path / 'merged.xlsx'
    book = Workbook()
    book.active.merge_cells('A1:B1')
    book.active['A1'] = 'merged'
    book.save(path)
    editor = editor_class().from_path(path)
    assert not editor.table.item(0, 1).flags() & __import__('PySide6.QtCore', fromlist=['Qt']).Qt.ItemIsEditable
    editor.table.item(1, 0).setText('=HYPERLINK("https://invalid.example", "text")')
    editor.save(path)
    result = load_workbook(path)
    assert 'A1:B1' in result.active.merged_cells
    assert result.active['A2'].value.startswith('=HYPERLINK')


def test_word_field_is_viewable_but_cannot_save(app, tmp_path):
    from docx.oxml import OxmlElement
    path = tmp_path / 'field.docx'
    doc = Document()
    paragraph = doc.add_paragraph('可查看')
    field = OxmlElement('w:fldSimple')
    field.set(__import__('docx.oxml.ns', fromlist=['qn']).qn('w:instr'), 'DATE')
    paragraph._p.append(field)
    doc.save(path)
    editor = editor_class().from_path(path)
    assert '可查看' in editor.text_edit.toPlainText()
    with pytest.raises(ValueError):
        editor.save(path)


def test_excel_view_does_not_materialize_empty_cells(app):
    editor = editor_class().blank('xlsx')
    assert len(editor.workbook.active._cells) == 0
    editor.table.item(0, 0).setText('5')
    assert len(editor.workbook.active._cells) == 1


def test_word_never_loads_external_image_resources(app):
    from PySide6.QtCore import QUrl
    from PySide6.QtGui import QTextDocument
    editor = editor_class().blank('docx')
    assert editor.text_edit.document().resource(QTextDocument.ImageResource, QUrl('https://example.invalid/image.png')) is None
    assert editor.text_edit.document().resource(QTextDocument.ImageResource, QUrl.fromLocalFile(__file__)) is None


@pytest.mark.parametrize('kind,extension', [('docx', 'xlsx'), ('xlsx', 'docx')])
def test_save_rejects_wrong_extension(app, tmp_path, kind, extension):
    editor = editor_class().blank(kind)
    with pytest.raises(ValueError):
        editor.save(tmp_path / ('wrong.' + extension))
    assert not list(tmp_path.iterdir())


def test_invalid_zip_is_value_error(app, tmp_path):
    path = tmp_path / 'bad.docx'
    path.write_bytes(b'not zip')
    with pytest.raises(ValueError):
        editor_class().from_path(path)


def test_word_styles_survive_edit_snapshot_and_unchanged_save(app, tmp_path):
    path = tmp_path / 'styles.docx'
    doc = Document()
    doc.add_paragraph('Heading', style='Heading 1')
    table = doc.add_table(rows=1, cols=1)
    table.style = 'Light Shading Accent 1'
    table.cell(0, 0).text = 'cell'
    doc.save(path)
    original = path.read_bytes()
    editor = editor_class().from_path(path)
    editor.save(path)
    assert path.read_bytes() == original
    editor.text_edit.moveCursor(QTextCursor.End)
    editor.text_edit.insertPlainText('extra')
    editor = editor_class().from_snapshot(editor.snapshot())
    editor.save(path)
    result = Document(path)
    assert result.paragraphs[0].style.style_id == 'Heading1'
    assert result.tables[0].style.style_id == 'LightShading-Accent1'


def test_word_paragraph_section_refuses_save(app, tmp_path):
    path = tmp_path / 'sections.docx'
    doc = Document()
    doc.add_paragraph('first')
    doc.add_section()
    doc.add_paragraph('second')
    doc.save(path)
    editor = editor_class().from_path(path)
    with pytest.raises(ValueError):
        editor.save(path)


@pytest.mark.parametrize('operation', ['save', 'snapshot'])
def test_excel_active_delegate_is_committed(app, tmp_path, operation):
    from PySide6.QtWidgets import QLineEdit
    from PySide6.QtTest import QTest
    from PySide6.QtCore import Qt
    editor = editor_class().blank('xlsx')
    editor.show()
    editor.table.setCurrentCell(0, 0)
    editor.table.editItem(editor.table.item(0, 0))
    app.processEvents()
    line = editor.table.findChild(QLineEdit)
    assert line is not None
    line.setText('123')
    if operation == 'save':
        path = tmp_path / 'active.xlsx'
        editor.save(path)
        assert load_workbook(path).active['A1'].value == 123
    else:
        restored = editor_class().from_snapshot(editor.snapshot())
        assert restored.workbook.active['A1'].value == 123
        assert not restored.undo_stack.isClean()
        restored.undo_stack.undo()
        assert restored.workbook.active['A1'].value is None
    editor.close()


def test_excel_rich_text_survives_save_undo_and_snapshot(app, tmp_path):
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    path = tmp_path / 'rich.xlsx'
    book = Workbook()
    book.active['A1'] = CellRichText('plain', TextBlock(InlineFont(b=True, color='FF0000'), 'bold'))
    book.save(path)
    editor = editor_class().from_path(path)
    assert isinstance(editor.workbook.active['A1'].value, CellRichText)
    editor.table.item(0, 1).setText('1')
    editor.undo_stack.undo()
    editor.undo_stack.redo()
    editor = editor_class().from_snapshot(editor.snapshot())
    editor.save(path)
    rich = load_workbook(path, rich_text=True).active['A1'].value
    assert isinstance(rich, CellRichText) and rich[1].font.b
    editor.table.item(0, 0).setText('replacement')
    editor.undo_stack.undo()
    assert isinstance(editor.workbook.active['A1'].value, CellRichText)


def _isolated_clipboard_check(check):
    """独立 offscreen Qt 进程使用内存剪贴板，不触碰 Windows 共享剪贴板。"""
    import os
    from pathlib import Path
    import subprocess
    import sys
    root = Path(__file__).resolve().parents[1]
    code = (
        'import runpy, sys; '
        'sys.path.insert(0, "tests"); '
        'from PySide6.QtWidgets import QApplication; '
        'app = QApplication([]); '
        'assert app.platformName() == "offscreen"; '
        'namespace = runpy.run_path("tests/test_office_editor.py"); '
        'namespace[sys.argv[1]](app)'
    )
    environment = dict(os.environ, QT_QPA_PLATFORM='offscreen')
    result = subprocess.run([sys.executable, '-c', code, check], cwd=root,
                            env=environment, capture_output=True, text=True, timeout=30)
    assert result.returncode == 0, result.stdout + result.stderr


def test_edit_action_excel_tsv_is_one_undo():
    _isolated_clipboard_check('_check_edit_action_excel_tsv_is_one_undo')


def _check_edit_action_excel_tsv_is_one_undo(app):
    editor = editor_class().blank('xlsx')
    editor.table.setCurrentCell(0, 0)
    app.clipboard().setText('12\tTRUE\n=A1\ttext')
    editor.edit_action('paste')
    assert editor.workbook.active['A1'].value == 12
    assert editor.workbook.active['B1'].value is True
    assert editor.workbook.active['A2'].value == '=A1'
    assert editor.undo_stack.count() == 1
    editor.undo_stack.undo()
    assert editor.workbook.active['A1'].value is None
    editor.undo_stack.redo()
    from PySide6.QtWidgets import QTableWidgetSelectionRange
    editor.table.setRangeSelected(QTableWidgetSelectionRange(0, 0, 1, 1), True)
    editor.edit_action('copy')
    assert app.clipboard().text() == '12\tTrue\n=A1\ttext'
    editor.edit_action('cut')
    assert editor.workbook.active['B2'].value is None
    editor.undo_stack.undo()
    assert editor.workbook.active['B2'].value == 'text'


def test_edit_action_word():
    _isolated_clipboard_check('_check_edit_action_word')


def _check_edit_action_word(app):
    editor = editor_class().blank('docx')
    editor.text_edit.insertPlainText('copy')
    editor.edit_action('selectAll')
    editor.edit_action('cut')
    assert editor.text_edit.toPlainText() == ''
    assert app.clipboard().text() == 'copy'
    assert editor.text_edit.canPaste()
    app.processEvents()
    editor.edit_action('paste')
    assert editor.text_edit.toPlainText() == 'copy'


def test_edit_action_excel_targets_active_delegate():
    _isolated_clipboard_check('_check_edit_action_excel_targets_active_delegate')


def _check_edit_action_excel_targets_active_delegate(app):
    from PySide6.QtWidgets import QLineEdit
    editor = editor_class().blank('xlsx')
    editor.show()
    editor.table.setCurrentCell(0, 0)
    editor.table.editItem(editor.table.item(0, 0))
    app.processEvents()
    line = editor.table.findChild(QLineEdit)
    line.setText('delegate')
    editor.edit_action('selectAll')
    editor.edit_action('cut')
    assert line.text() == ''
    assert app.clipboard().text() == 'delegate'
    editor.edit_action('paste')
    assert line.text() == 'delegate'
    editor.snapshot()
    assert editor.workbook.active['A1'].value == 'delegate'
    editor.close()


def test_excel_accepting_unchanged_rich_text_preserves_runs(app, tmp_path):
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from PySide6.QtWidgets import QLineEdit
    path = tmp_path / 'unchanged-rich.xlsx'
    book = Workbook()
    book.active['A1'] = CellRichText('plain', TextBlock(InlineFont(b=True), 'bold'))
    book.save(path)
    editor = editor_class().from_path(path)
    editor.show()
    editor.table.setCurrentCell(0, 0)
    editor.table.editItem(editor.table.item(0, 0))
    app.processEvents()
    editor.snapshot()
    assert isinstance(editor.workbook.active['A1'].value, CellRichText)
    assert not editor.undo_stack.canUndo()
    editor.close()


@pytest.mark.parametrize('change', [
    {'office_format': 'xlsm'},
    {'states': []},
    {'source': 'not base64!'},
    {'office_format': 'xlsx', 'states': [{'book': 'bad!', 'sheet': 0}]},
])
def test_corrupt_snapshot_rejected(app, change):
    snapshot = editor_class().blank('docx').snapshot()
    snapshot.update(change)
    with pytest.raises(ValueError):
        editor_class().from_snapshot(snapshot)
