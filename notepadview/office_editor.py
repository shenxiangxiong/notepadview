"""基础 Office 编辑；复杂结构只读保存保护，公式仅作为文本处理。"""
import base64
from copy import deepcopy
from datetime import date, datetime
from html import escape
import io
import json
import os
from pathlib import Path
import re
import uuid
import zipfile
import xml.etree.ElementTree as ET

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import Workbook, load_workbook
from PySide6.QtCore import Qt, Signal, QUrl
from PySide6.QtGui import (QColor, QFont, QTextCharFormat, QTextCursor, QTextDocument,
                          QTextTable, QTextFormat, QUndoCommand, QUndoStack)
from PySide6.QtWidgets import (QApplication, QLineEdit, QAbstractItemDelegate, QWidget, QVBoxLayout, QToolBar, QTextEdit, QTableWidget,
    QTableWidgetItem, QComboBox, QInputDialog, QColorDialog, QFontComboBox, QSpinBox)

_MAX_PACKAGE = 64 * 1024 * 1024
_MAX_EXPANDED = 128 * 1024 * 1024
_MAX_CELLS = 200000


def _package(data, kind):
    if len(data) > _MAX_PACKAGE:
        raise ValueError('Office 文件超过读取限制。')
    unsafe = []
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 4096 or sum(i.file_size for i in entries) > _MAX_EXPANDED:
                raise ValueError('Office ZIP 解压尺寸超过限制。')
            for entry in entries:
                name = entry.filename
                if entry.flag_bits & 1 or entry.file_size > 32 * 1024 * 1024:
                    raise ValueError('Office ZIP 包含加密或过大部件。')
                if entry.file_size > max(1, entry.compress_size) * 1000:
                    raise ValueError('Office ZIP 压缩比超过限制。')
                low = name.lower()
                if kind == 'docx':
                    if any(s in low for s in ('footnotes', 'endnotes', 'header', 'footer', 'comments', 'embeddings/', 'charts/', 'diagrams/', 'vba')):
                        unsafe.append(name)
                elif any(s in low for s in ('vbaproject', 'externallinks/', 'drawings/', 'charts/', 'embeddings/', 'pivot', 'slicer', 'activex', 'ctrlprops', 'connections', 'querytables', 'customxml/', 'threadedcomments', 'richdata', 'persons/')):
                    unsafe.append(name)
                if name.endswith(('.xml', '.rels')):
                    content = archive.read(entry)
                    if b'<!DOCTYPE' in content.upper() or b'<!ENTITY' in content.upper():
                        raise ValueError('Office XML 不允许实体声明。')
                    root = ET.fromstring(content)
                    if name.endswith('.rels') and any(e.get('TargetMode') == 'External' for e in root):
                        unsafe.append('外部关系')
                    if kind == 'docx' and name == 'word/document.xml':
                        unsupported = {'ins', 'del', 'moveFrom', 'moveTo', 'fldSimple', 'fldChar', 'instrText', 'object', 'altChunk', 'hyperlink', 'sdt', 'numPr', 'oMath', 'oMathPara', 'anchor', 'txbxContent', 'gridSpan', 'vMerge', 'bookmarkStart', 'commentRangeStart', 'footnoteReference', 'endnoteReference', 'pict'}
                        if any(e.tag.split('}')[-1] == 'sectPr' for p in root.iter() if p.tag.split('}')[-1] == 'pPr' for e in p):
                            unsafe.append('分节设置')
                        if any(e.tag.split('}')[-1] in unsupported for e in root.iter()):
                            unsafe.append('复杂 Word 结构')
                        if any(sum(1 for node in table.iter() if node.tag.split('}')[-1] == 'tbl') > 1
                               for table in root.iter() if table.tag.split('}')[-1] == 'tbl'):
                            unsafe.append('嵌套表格')
                    if kind == 'xlsx' and any(e.tag.split('}')[-1] == 'extLst' for e in root.iter()):
                        unsafe.append('Excel 扩展对象')
            required = 'word/document.xml' if kind == 'docx' else 'xl/workbook.xml'
            if required not in archive.namelist():
                raise ValueError('不是有效的 Office 文档。')
    except (zipfile.BadZipFile, ET.ParseError, RuntimeError, KeyError) as error:
        raise ValueError('Office 文件无效或损坏。') from error
    return unsafe


def _book_bytes(book):
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


class _EmbeddedDocument(QTextDocument):
    def loadResource(self, resource_type, name):
        if name.scheme() != 'data':
            return None
        return super().loadResource(resource_type, name)


class _Change(QUndoCommand):
    def __init__(self, owner, before, after):
        super().__init__('Office 编辑')
        self.owner, self.before, self.after = owner, before, after

    def undo(self):
        self.owner._restore(self.before)

    def redo(self):
        self.owner._restore(self.after)


class OfficeEditor(QWidget):
    changed = Signal()

    def __init__(self, office_format, parent=None):
        super().__init__(parent)
        if office_format not in ('docx', 'xlsx'):
            raise ValueError('仅支持 docx 和 xlsx。')
        self.office_format = office_format
        self.undo_stack = QUndoStack(self)
        self._busy = True
        self._source = b''
        self._unsafe = []
        self._states = []
        self._sheet = 0
        self._batch = False
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        self.toolbar = QToolBar('Office 编辑', self)
        layout.addWidget(self.toolbar)
        self.toolbar.addAction('撤销', self.undo_stack.undo)
        self.toolbar.addAction('重做', self.undo_stack.redo)
        if office_format == 'docx':
            self.text_edit = QTextEdit(self)
            self.text_edit.setDocument(_EmbeddedDocument(self.text_edit))
            self.text_edit.setAcceptRichText(False)
            self.text_edit.setUndoRedoEnabled(False)
            layout.addWidget(self.text_edit)
            for title, attribute in [('粗体', 'bold'), ('斜体', 'italic'), ('下划线', 'underline')]:
                action = self.toolbar.addAction(title)
                action.setCheckable(True)
                action.toggled.connect(lambda enabled, key=attribute: self._format(key, enabled))
            fonts = QFontComboBox(self.toolbar)
            fonts.currentFontChanged.connect(lambda font: self._format('family', font.family()))
            self.toolbar.addWidget(fonts)
            sizes = QSpinBox(self.toolbar)
            sizes.setRange(1, 200)
            sizes.setValue(11)
            sizes.setToolTip('字号')
            sizes.valueChanged.connect(lambda size: self._format('size', size))
            self.toolbar.addWidget(sizes)
            self.toolbar.addAction('文字颜色', self._choose_color)
            for title, alignment in [('左对齐', Qt.AlignLeft), ('居中', Qt.AlignHCenter), ('右对齐', Qt.AlignRight)]:
                self.toolbar.addAction(title, lambda checked=False, value=alignment: self.text_edit.setAlignment(value))
            self.toolbar.addAction('插入表格', self._insert_table)
            self.text_edit.textChanged.connect(self._record)
        else:
            self.workbook = Workbook()
            self.sheet_selector = QComboBox(self)
            self.sheet_selector.currentIndexChanged.connect(self._select_sheet)
            self.toolbar.addWidget(self.sheet_selector)
            self.toolbar.addAction('添加工作表', lambda: self._sheet_dialog(False))
            self.toolbar.addAction('重命名工作表', lambda: self._sheet_dialog(True))
            self.toolbar.addAction('删除工作表', self.delete_sheet)
            self.table = QTableWidget(self)
            layout.addWidget(self.table)
            self.table.itemChanged.connect(self._cell_changed)
        self.undo_stack.indexChanged.connect(self._notify_change)

    def set_view_only(self):
        self.view_only = True
        for action in self.toolbar.actions():
            action.setVisible(False)
        if self.office_format == 'docx':
            self.text_edit.setReadOnly(True)
            self.toolbar.hide()
        else:
            for action in self.toolbar.actions():
                if self.toolbar.widgetForAction(action) is self.sheet_selector:
                    action.setVisible(True)
            self.table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.undo_stack.setClean()

    def _notify_change(self, index):
        self.changed.emit()

    @classmethod
    def blank(cls, format, parent=None):
        editor = cls(format, parent)
        if format == 'xlsx':
            editor._render_sheet()
        editor._finish_load()
        return editor

    @classmethod
    def from_path(cls, path, parent=None):
        path = Path(path)
        kind = path.suffix.lower().lstrip('.')
        if kind not in ('docx', 'xlsx'):
            raise ValueError('仅支持 docx 和 xlsx。')
        if path.stat().st_size > _MAX_PACKAGE:
            raise ValueError('Office 文件超过读取限制。')
        data = path.read_bytes()
        unsafe = _package(data, kind)
        editor = cls(kind, parent)
        editor._source, editor._unsafe = data, unsafe
        try:
            if kind == 'docx':
                editor._load_word(data)
            else:
                editor.workbook = load_workbook(io.BytesIO(data), data_only=False, keep_links=False, rich_text=True)
                editor._render_sheet()
        except Exception as error:
            editor.deleteLater()
            raise ValueError('无法读取 Office 文档。') from error
        editor._finish_load()
        return editor

    def _finish_load(self):
        self._busy = False
        self._state = self._capture()
        self._states = [deepcopy(self._state)]
        self.undo_stack.setClean()

    def _capture(self):
        if self.office_format == 'docx':
            document = self.text_edit.document()
            metadata = []
            block = document.begin()
            while block.isValid():
                value = block.blockFormat().property(QTextFormat.UserProperty)
                if value:
                    metadata.append([block.position(), value])
                block = block.next()
            tables = []
            iterator = document.rootFrame().begin()
            while not iterator.atEnd():
                frame = iterator.currentFrame()
                if isinstance(frame, QTextTable):
                    tables.append([frame.firstPosition(), frame.format().property(QTextFormat.UserProperty)])
                iterator += 1
            return {'html': document.toHtml(), 'paragraphs': metadata, 'tables': tables}
        return {'book': base64.b64encode(_book_bytes(self.workbook)).decode('ascii'), 'sheet': self._sheet}

    def _restore(self, state):
        self._busy = True
        try:
            if self.office_format == 'docx':
                position = self.text_edit.textCursor().position()
                self.text_edit.setHtml(state['html'])
                document = self.text_edit.document()
                for offset, value in state.get('paragraphs', []):
                    mark = QTextCursor(document.findBlock(offset))
                    fmt = mark.blockFormat()
                    fmt.setProperty(QTextFormat.UserProperty, value)
                    mark.setBlockFormat(fmt)
                for offset, value in state.get('tables', []):
                    mark = QTextCursor(document)
                    mark.setPosition(offset)
                    table = mark.currentTable()
                    if table is not None and value:
                        fmt = table.format()
                        fmt.setProperty(QTextFormat.UserProperty, value)
                        table.setFormat(fmt)
                cursor = self.text_edit.textCursor()
                cursor.setPosition(min(position, self.text_edit.document().characterCount() - 1))
                self.text_edit.setTextCursor(cursor)
            else:
                self.workbook = load_workbook(io.BytesIO(base64.b64decode(state['book'])), data_only=False, keep_links=False, rich_text=True)
                self._sheet = min(state['sheet'], len(self.workbook.worksheets) - 1)
                self._render_sheet()
            self._state = deepcopy(state)
        finally:
            self._busy = False

    def _record(self):
        if self._busy or self._batch:
            return
        state = self._capture()
        if state == self._state:
            return
        index = self.undo_stack.index()
        self._states = self._states[:index + 1] + [deepcopy(state)]
        self.undo_stack.push(_Change(self, deepcopy(self._state), state))

    def snapshot(self):
        self._commit_delegate()
        return {'version': 1, 'office_format': self.office_format,
                'source': base64.b64encode(self._source).decode('ascii'),
                'unsafe': list(self._unsafe), 'states': deepcopy(self._states),
                'index': self.undo_stack.index(), 'clean_index': self.undo_stack.cleanIndex()}

    @classmethod
    def from_snapshot(cls, snapshot, parent=None):
        try:
            if snapshot['version'] != 1:
                raise ValueError('不支持的 Office 草稿版本。')
            editor = cls(snapshot['office_format'], parent)
            editor._source = base64.b64decode(snapshot['source'], validate=True)
            editor._unsafe = list(snapshot.get('unsafe', []))
            if editor._source:
                editor._unsafe.extend(_package(editor._source, editor.office_format))
            states = deepcopy(snapshot['states'])
            if not states or len(json.dumps(states)) > _MAX_EXPANDED * 2:
                raise ValueError('Office 草稿超过限制。')
            for state in states:
                if editor.office_format == 'xlsx':
                    _package(base64.b64decode(state['book'], validate=True), 'xlsx')
            editor._restore(states[0])
            editor._states = states
            for before, after in zip(states, states[1:]):
                editor.undo_stack.push(_Change(editor, before, after))
            clean = snapshot['clean_index']
            if 0 <= clean < len(states):
                editor.undo_stack.setIndex(clean)
                editor.undo_stack.setClean()
            else:
                editor.undo_stack.resetClean()
            editor.undo_stack.setIndex(snapshot['index'])
            editor._busy = False
            return editor
        except Exception as error:
            raise ValueError('Office 草稿无效。') from error

    def _format(self, key, value):
        fmt = QTextCharFormat()
        if key == 'bold':
            fmt.setFontWeight(QFont.Bold if value else QFont.Normal)
        elif key == 'italic':
            fmt.setFontItalic(value)
        elif key == 'underline':
            fmt.setFontUnderline(value)
        elif key == 'family':
            fmt.setFontFamilies([value])
        elif key == 'size':
            fmt.setFontPointSize(value)
        elif key == 'color':
            fmt.setForeground(value)
        self.text_edit.mergeCurrentCharFormat(fmt)

    def _choose_color(self):
        color = QColorDialog.getColor(parent=self)
        if color.isValid():
            self._format('color', color)

    def _insert_table(self):
        rows, ok = QInputDialog.getInt(self, '插入表格', '行数', 2, 1, 100)
        if ok:
            cols, ok = QInputDialog.getInt(self, '插入表格', '列数', 2, 1, 30)
            if ok:
                self.text_edit.textCursor().insertTable(rows, cols)

    def _load_word(self, data):
        doc = Document(io.BytesIO(data))
        from docx.text.paragraph import Paragraph
        from docx.table import Table
        paragraph_properties = {}
        table_properties = []
        from lxml.etree import tostring
        def paragraph(p):
            marker = 'nvparagraph' + uuid.uuid4().hex
            paragraph_properties[marker] = tostring(p._p.pPr, encoding='unicode') if p._p.pPr is not None else ''
            align = {0: 'left', 1: 'center', 2: 'right', 3: 'justify'}.get(p.alignment, 'left')
            content = []
            for run in p.runs:
                font = run.font
                style_font = run.style.font
                pfont = p.style.font
                def inherited(attr):
                    for f in (font, style_font, pfont):
                        value = getattr(f, attr)
                        if value is not None:
                            return value
                    return None
                styles = []
                for attr, css, val in [('bold', 'font-weight', 'bold'), ('italic', 'font-style', 'italic'), ('underline', 'text-decoration', 'underline')]:
                    if inherited(attr):
                        styles.append(f'{css}:{val}')
                if inherited('name'):
                    styles.append('font-family:' + escape(inherited('name'), quote=True))
                if inherited('size'):
                    styles.append(f'font-size:{inherited("size").pt}pt')
                if font.color.rgb:
                    styles.append('color:#' + str(font.color.rgb))
                content.append('<span style="' + ';'.join(styles) + '">' + escape(run.text).replace('\n', '<br/>').replace('\t', '&#9;') + '</span>')
                for drawing in run._r.iter():
                    if drawing.tag.split('}')[-1] == 'blip':
                        rid = drawing.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                        if rid in doc.part.related_parts:
                            part = doc.part.related_parts[rid]
                            encoded = base64.b64encode(part.blob).decode('ascii')
                            extent = next((e for e in run._r.iter() if e.tag.split('}')[-1] == 'extent'), None)
                            dimensions = ''
                            if extent is not None:
                                dimensions = f' width="{int(extent.get("cx", 914400)) / 9525}" height="{int(extent.get("cy", 914400)) / 9525}"'
                            content.append(f'<img src="data:{part.content_type};base64,{encoded}"{dimensions}/>')
            return f'<p align="{align}" style="white-space:pre-wrap">' + marker + ''.join(content) + '</p>'
        html = []
        for child in doc.element.body:
            tag = child.tag.split('}')[-1]
            if tag == 'p':
                html.append(paragraph(Paragraph(child, doc)))
            elif tag == 'tbl':
                table = Table(child, doc)
                table_properties.append(tostring(child.tblPr, encoding='unicode'))
                html.append('<table border="1" cellspacing="0">')
                for row in table.rows:
                    html.append('<tr>')
                    for cell in row.cells:
                        html.append('<td>' + ''.join(paragraph(p) for p in cell.paragraphs) + '</td>')
                    html.append('</tr>')
                html.append('</table>')
        self.text_edit.setHtml(''.join(html))
        document = self.text_edit.document()
        for marker, value in paragraph_properties.items():
            cursor = document.find(marker)
            cursor.removeSelectedText()
            fmt = cursor.blockFormat()
            fmt.setProperty(QTextFormat.UserProperty, value)
            cursor.setBlockFormat(fmt)
        iterator = document.rootFrame().begin()
        index = 0
        while not iterator.atEnd():
            frame = iterator.currentFrame()
            if isinstance(frame, QTextTable):
                fmt = frame.format()
                fmt.setProperty(QTextFormat.UserProperty, table_properties[index])
                frame.setFormat(fmt)
                index += 1
            iterator += 1

    def _word_bytes(self):
        doc = Document(io.BytesIO(self._source)) if self._source else Document()
        for child in list(doc.element.body):
            if child.tag.split('}')[-1] != 'sectPr':
                doc.element.body.remove(child)
        qt_doc = self.text_edit.document()
        def add_block(block, target):
            p = target.add_paragraph()
            properties = block.blockFormat().property(QTextFormat.UserProperty)
            if properties:
                from docx.oxml import parse_xml
                p._p.insert(0, parse_xml(properties))
            alignment = block.blockFormat().alignment()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if alignment & Qt.AlignHCenter else WD_ALIGN_PARAGRAPH.RIGHT if alignment & Qt.AlignRight else WD_ALIGN_PARAGRAPH.JUSTIFY if alignment & Qt.AlignJustify else WD_ALIGN_PARAGRAPH.LEFT
            it = block.begin()
            while not it.atEnd():
                fragment = it.fragment()
                it += 1
                if not fragment.isValid():
                    continue
                fmt = fragment.charFormat()
                if fmt.isImageFormat():
                    image_fmt = fmt.toImageFormat()
                    name = image_fmt.name()
                    if name.startswith('data:'):
                        payload = base64.b64decode(name.split(',', 1)[1])
                    else:
                        raise ValueError('无法保存未嵌入的图片。')
                    kwargs = {}
                    if image_fmt.width() > 0:
                        kwargs['width'] = Inches(image_fmt.width() / 96)
                    if image_fmt.height() > 0:
                        kwargs['height'] = Inches(image_fmt.height() / 96)
                    p.add_run().add_picture(io.BytesIO(payload), **kwargs)
                else:
                    run = p.add_run(fragment.text().replace('\u2028', '\n'))
                    run.bold = fmt.fontWeight() >= QFont.Bold
                    run.italic, run.underline = fmt.fontItalic(), fmt.fontUnderline()
                    if fmt.fontFamilies():
                        run.font.name = fmt.fontFamilies()[0]
                    if fmt.fontPointSize() > 0:
                        run.font.size = Pt(fmt.fontPointSize())
                    if fmt.hasProperty(QTextCharFormat.ForegroundBrush):
                        color = fmt.foreground().color()
                        run.font.color.rgb = RGBColor(color.red(), color.green(), color.blue())
            return p
        iterator = qt_doc.rootFrame().begin()
        while not iterator.atEnd():
            frame = iterator.currentFrame()
            if isinstance(frame, QTextTable):
                table = doc.add_table(rows=frame.rows(), cols=frame.columns())
                properties = frame.format().property(QTextFormat.UserProperty)
                if properties:
                    from docx.oxml import parse_xml
                    table._tbl.replace(table._tbl.tblPr, parse_xml(properties))
                for row in range(frame.rows()):
                    for col in range(frame.columns()):
                        cell = frame.cellAt(row, col)
                        target = table.cell(row, col)
                        target._element.remove(target.paragraphs[0]._element)
                        block = cell.firstCursorPosition().block()
                        end = cell.lastCursorPosition().position()
                        while block.isValid() and block.position() <= end:
                            add_block(block, target)
                            block = block.next()
            elif iterator.currentBlock().isValid():
                add_block(iterator.currentBlock(), doc)
            iterator += 1
        stream = io.BytesIO()
        doc.save(stream)
        return stream.getvalue()

    def _render_sheet(self):
        previous = self._busy
        self._busy = True
        try:
            self.sheet_selector.blockSignals(True)
            self.sheet_selector.clear()
            self.sheet_selector.addItems(self.workbook.sheetnames)
            self.sheet_selector.setCurrentIndex(self._sheet)
            self.sheet_selector.blockSignals(False)
            sheet = self.workbook.worksheets[self._sheet]
            rows, cols = max(100, sheet.max_row + 20), max(26, sheet.max_column + 5)
            if rows * cols > _MAX_CELLS:
                raise ValueError('工作表单元格范围超过编辑限制。')
            self.table.clearSpans()
            self.table.clear()
            self.table.setRowCount(rows)
            self.table.setColumnCount(cols)
            from openpyxl.utils import get_column_letter
            self.table.setHorizontalHeaderLabels([get_column_letter(c + 1) for c in range(cols)])
            for row in range(rows):
                for col in range(cols):
                    cell = sheet._cells.get((row + 1, col + 1))
                    value = cell.value if cell is not None else None
                    text = '' if value is None else str(value)
                    item = QTableWidgetItem(text)
                    from openpyxl.cell.cell import MergedCell
                    if isinstance(cell, MergedCell):
                        item.setFlags(item.flags() & ~Qt.ItemIsEditable)
                    self.table.setItem(row, col, item)
            for merged in sheet.merged_cells.ranges:
                self.table.setSpan(merged.min_row - 1, merged.min_col - 1,
                                   merged.max_row - merged.min_row + 1,
                                   merged.max_col - merged.min_col + 1)
        finally:
            self._busy = previous

    def _select_sheet(self, index):
        if not self._busy and index >= 0:
            self._sheet = index
            self._render_sheet()

    def _cell_changed(self, item):
        if self._busy:
            return
        sheet = self.workbook.worksheets[self._sheet]
        cell = sheet.cell(item.row() + 1, item.column() + 1)
        text = item.text()
        if text == '':
            value = None
        elif text.startswith('='):
            value = text
        elif text.lower() in ('true', 'false'):
            value = text.lower() == 'true'
        elif re.fullmatch(r'[+-]?\d+', text) and not re.fullmatch(r'[+-]?0\d+', text):
            value = int(text)
        elif re.fullmatch(r'[+-]?(?:\d+\.\d*|\d*\.\d+)(?:[eE][+-]?\d+)?', text):
            value = float(text)
        elif isinstance(cell.value, (date, datetime)):
            try:
                value = datetime.fromisoformat(text)
            except ValueError:
                value = text
        else:
            value = text
        cell.value = value
        self._record()

    def _sheet_dialog(self, rename):
        name, ok = QInputDialog.getText(self, '工作表', '工作表名称')
        if ok and name:
            try:
                self.rename_sheet(name) if rename else self.add_sheet(name)
            except ValueError:
                from PySide6.QtWidgets import QMessageBox
                QMessageBox.warning(self, '工作表', '工作表名称无效或重复。')

    def _validate_name(self, name):
        if not isinstance(name, str) or not name.strip() or len(name) > 31 or re.search(r'[\\/*?:\[\]]', name):
            raise ValueError('工作表名称无效。')
        if name.casefold() in [n.casefold() for n in self.workbook.sheetnames]:
            raise ValueError('工作表名称已存在。')

    def add_sheet(self, name=None):
        if name is not None:
            self._validate_name(name)
        sheet = self.workbook.create_sheet(name)
        self._sheet = self.workbook.index(sheet)
        self._record()

    def rename_sheet(self, name):
        if name == self.workbook.worksheets[self._sheet].title:
            return
        self._validate_name(name)
        self.workbook.worksheets[self._sheet].title = name
        self._record()

    def delete_sheet(self):
        if len(self.workbook.worksheets) <= 1:
            return
        self.workbook.remove(self.workbook.worksheets[self._sheet])
        self._sheet = min(self._sheet, len(self.workbook.worksheets) - 1)
        self._record()

    def _active_delegate(self):
        if self.office_format != 'xlsx':
            return None
        return next((line for line in self.table.findChildren(QLineEdit)
                     if line.isVisible()), None)

    def _commit_delegate(self):
        line = self._active_delegate()
        if line is None:
            return
        from shiboken6 import isValid
        self.table.itemDelegate().setModelData(line, self.table.model(), self.table.currentIndex())
        if isValid(line):
            self.table.closeEditor(line, QAbstractItemDelegate.NoHint)

    def edit_action(self, action: str) -> None:
        if getattr(self, 'view_only', False) and action not in ('copy', 'selectAll'):
            return
        if action not in ('copy', 'cut', 'paste', 'selectAll'):
            raise ValueError('不支持的编辑操作。')
        if self.office_format == 'docx':
            getattr(self.text_edit, action)()
            return
        line = self._active_delegate()
        if line is not None:
            getattr(line, action)()
            return
        if action == 'selectAll':
            self.table.selectAll()
            return
        indexes = self.table.selectedIndexes()
        if action in ('copy', 'cut'):
            if not indexes:
                return
            top, bottom = min(i.row() for i in indexes), max(i.row() for i in indexes)
            left, right = min(i.column() for i in indexes), max(i.column() for i in indexes)
            text = '\n'.join('\t'.join(self.table.item(row, col).text()
                           for col in range(left, right + 1)) for row in range(top, bottom + 1))
            QApplication.clipboard().setText(text)
            if action == 'copy':
                return
            changes = [(i.row(), i.column(), '') for i in indexes]
        else:
            row, col = max(0, self.table.currentRow()), max(0, self.table.currentColumn())
            changes = [(row + dr, col + dc, text)
                       for dr, values in enumerate(QApplication.clipboard().text().splitlines())
                       for dc, text in enumerate(values.split('\t'))]
            if any(r >= self.table.rowCount() or c >= self.table.columnCount() for r, c, _ in changes):
                raise ValueError('粘贴范围超过工作表编辑区域。')
        self._batch = True
        try:
            for row, col, text in changes:
                item = self.table.item(row, col)
                if item.flags() & Qt.ItemIsEditable:
                    item.setText(text)
        finally:
            self._batch = False
        self._record()

    def save(self, path):
        if getattr(self, 'view_only', False):
            raise ValueError('Office 文档仅支持查看。')
        self._commit_delegate()
        if self._unsafe:
            raise ValueError('文档包含无法无损保留的结构，已拒绝保存：' + ', '.join(sorted(set(self._unsafe))))
        path = Path(path)
        if path.suffix.lower() != '.' + self.office_format:
            raise ValueError('保存扩展名必须与 Office 格式一致。')
        if self.office_format == 'docx':
            payload = self._source if self._source and self._state == self._states[0] else self._word_bytes()
        else:
            payload = _book_bytes(self.workbook)
        temporary = path.with_name('.' + path.name + '.' + uuid.uuid4().hex + '.writing')
        try:
            with temporary.open('xb') as stream:
                stream.write(payload)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary, path)
        finally:
            temporary.unlink(missing_ok=True)
        self.undo_stack.setClean()
